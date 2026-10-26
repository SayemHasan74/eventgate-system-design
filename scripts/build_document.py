from pathlib import Path
import json,re
from collections import defaultdict
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.section import WD_SECTION_START,WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'deliverables';OUT.mkdir(exist_ok=True)
models=json.loads((ROOT/'design/schema-summary.json').read_text())
api=json.loads((ROOT/'design/api-catalogue.json').read_text())
links=json.loads((ROOT/'design/share-links.json').read_text())
doc=Document()
for s in doc.sections:
 s.page_width=Inches(8.27);s.page_height=Inches(11.69);s.top_margin=Inches(.7);s.bottom_margin=Inches(.65);s.left_margin=Inches(.7);s.right_margin=Inches(.7)
for n in ['Normal','Body Text','Title','Subtitle','Heading 1','Heading 2','Heading 3','Caption']:
 st=doc.styles[n];st.font.name='Arial';st.font.color.rgb=RGBColor(0,0,0)
 st.paragraph_format.space_after=Pt(6)
doc.styles['Normal'].font.size=Pt(10.5)
doc.styles['Normal'].paragraph_format.line_spacing=1.12
for n,size in [('Title',26),('Subtitle',13),('Heading 1',17),('Heading 2',13),('Heading 3',11)]:
 doc.styles[n].font.size=Pt(size);doc.styles[n].paragraph_format.keep_with_next=True
 if n!='Subtitle':doc.styles[n].font.bold=True
doc.styles['Caption'].font.size=Pt(9)
doc.styles.add_style('Code Example',1)
code_style=doc.styles['Code Example'];code_style.font.name='Courier New';code_style.font.size=Pt(8.4);code_style.paragraph_format.line_spacing=1;code_style.paragraph_format.space_after=Pt(0)
for st in [doc.styles['Title'],doc.styles['Subtitle']]:
 pr=st.element.find(qn('w:pPr'))
 if pr is not None:
  for b in list(pr.findall(qn('w:pBdr'))):pr.remove(b)
# A long technical specification needs page numbers for referencing contracts.
footer=doc.sections[0].footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=footer.add_run('EventGate  |  ');r.font.size=Pt(8)
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)
doc.core_properties.title='EventGate Backend System Design';doc.core_properties.author='Sayem Hasan';doc.core_properties.subject='Backend architecture database relationships and API specifications'

def p(s='',style=None):return doc.add_paragraph(s,style)
def h(s,level=1):return doc.add_heading(s,level)
def label(name,text):
 q=p();q.add_run(name+' ').bold=True;q.add_run(text);return q
def link(label_text,url):
 q=p();el=OxmlElement('w:hyperlink');el.set(qn('r:id'),q.part.relate_to(url,RT.HYPERLINK,is_external=True));run=OxmlElement('w:r');pr=OxmlElement('w:rPr');co=OxmlElement('w:color');co.set(qn('w:val'),'174D73');pr.append(co);run.append(pr);t=OxmlElement('w:t');t.text=label_text;run.append(t);el.append(run);q._p.append(el);return q
def table(headers,rows,widths=None,size=9):
 t=doc.add_table(rows=1,cols=len(headers));t.autofit=False;t.style='Table Grid'
 if widths:
  for col,w in zip(t.columns,widths):col.width=Inches(w)
  for cell,w in zip(t.rows[0].cells,widths):cell.width=Inches(w)
 for i,x in enumerate(headers):t.rows[0].cells[i].text=str(x)
 repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
 for row in rows:
  cells=t.add_row().cells
  for i,x in enumerate(row):cells[i].text=str(x)
 for ri,row in enumerate(t.rows):
  trpr=row._tr.get_or_add_trPr();trpr.append(OxmlElement('w:cantSplit'))
  for ci,cell in enumerate(row.cells):
   if widths:cell.width=Inches(widths[ci])
   tcpr=cell._tc.get_or_add_tcPr();mar=OxmlElement('w:tcMar')
   for edge in ['top','left','bottom','right']:
    e=OxmlElement('w:'+edge);e.set(qn('w:w'),'70');e.set(qn('w:type'),'dxa');mar.append(e)
   tcpr.append(mar)
   if ri==0:
    shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E8EDF1');tcpr.append(shade)
   for para in cell.paragraphs:
    para.paragraph_format.space_after=Pt(2);para.paragraph_format.line_spacing=1.05
    for run in para.runs:run.font.size=Pt(size);run.bold=ri==0
 p();return t
def bullets(items):
 for item in items:p(item,'List Bullet')
def numbered(items):
 for i,item in enumerate(items,1):p(f'{i}. {item}')
def section(landscape=False):
 s=doc.add_section(WD_SECTION_START.NEW_PAGE)
 s.orientation=WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
 s.page_width=Inches(11.69 if landscape else 8.27);s.page_height=Inches(8.27 if landscape else 11.69)
 s.left_margin=Inches(.55 if landscape else .7);s.right_margin=s.left_margin
 s.top_margin=Inches(.45 if landscape else .7);s.bottom_margin=Inches(.45 if landscape else .65)
 return s
def diagram(file,caption):
 para=p();para.paragraph_format.space_after=Pt(2);
 from PIL import Image
 im=Image.open(ROOT/'diagrams'/file);width=min(10.5,6.35*im.width/im.height)
 para.add_run().add_picture(str(ROOT/'diagrams'/file),width=Inches(width))
 p(caption,'Caption')
def code(obj):
 text=obj if isinstance(obj,str) else json.dumps(obj,indent=2,ensure_ascii=True)
 lines=text.splitlines()
 for idx,line in enumerate(lines):
  q=p(line,'Code Example');q.paragraph_format.keep_with_next=idx<len(lines)-1

def err(message,error_code,detail=None,field=None):
 e={'code':error_code,'message':detail or message}
 if field:e['field']=field
 return {'success':False,'message':message,'errors':[e]}
def ok(message,data):return {'success':True,'message':message,'data':data}
UID='11111111-1111-4111-8111-111111111111';EID='22222222-2222-4222-8222-222222222222';TID='33333333-3333-4333-8333-333333333333';OID='44444444-4444-4444-8444-444444444444';PID='55555555-5555-4555-8555-555555555555';TKID='66666666-6666-4666-8666-666666666666';RID='77777777-7777-4777-8777-777777777777'

p('EventGate Backend System Design','Title')
p('Architecture database design API specifications and system workflows','Subtitle')
p('Prepared by Sayem Hasan\n26 October 2026')
p('EventGate manages the full journey from event discovery and limited ticket inventory to verified payments, QR admission and refunds. This design explains how the existing project protects stock, money and admission state while serving attendees, organizers and administrators.')
p('The backend is a modular Express and TypeScript application using Prisma and PostgreSQL. A browser backend for frontend manages cookie sessions through Redis. Durable database jobs handle reservation expiry, payment verification, refunds and notifications. PostgreSQL remains the source of truth for inventory and financial outcomes.')
label('Project references','Existing EventGate application and backend source')
link('EventGate application','https://eventgate-frontend-lovat.vercel.app/')
link('EventGate backend repository','https://github.com/SayemHasan74/eventgate-api')
link('Open the completed ERD in Draw.io',links['erd'])
link('Open the SDA workflow and architecture diagrams',links['workflow'])
h('Document map',2)
table(['Section','Content'],[
 ('1','Project purpose target users eight core features and roles'),('2','Backend architecture and responsibilities'),('3','SDA workflow and transaction explanation'),('4','Database ERD entity dictionary and relationships'),('5','API conventions and complete operation catalogue'),('6','Detailed request and response designs for ten APIs'),('7','State transitions reliability security and validation checks')],[.7,6.1],10)

h('1 Project requirements')
label('Project name','EventGate')
label('Target users','People booking events in Bangladesh, organizers publishing and operating events, and administrators overseeing users, payments and operational exceptions.')
label('Main purpose','Offer a reliable event ticketing system where available stock cannot be oversold, payment is confirmed by the provider, each issued pass admits one attendee once, and refund requests remain traceable.')
h('Eight core features',2)
table(['Feature','User value and backend responsibility'],[
 ('1 Identity and account management','Registration, password and Google sign in, session rotation, profile settings, verification and role enforcement.'),
 ('2 Event discovery and personalization','Search and filter published events; read event facts and fresh availability; save events, choose preferences and request stock alerts. Optional AI assists discovery without modifying bookings.'),
 ('3 Organizer event lifecycle','Create drafts, manage ticket tiers, check publication readiness, publish valid events and process cancellations.'),
 ('4 Inventory reservation and order history','Reserve 1 to 10 tickets in one tier for 15 minutes; use authoritative stock checks, idempotency and price snapshots.'),
 ('5 Verified checkout and payment history','Create hosted SSLCommerz checkout, record attempts and callbacks, verify captures and reconcile uncertain outcomes.'),
 ('6 Private tickets and QR admission','Issue tickets only after verified fulfillment; protect QR retrieval; validate event, time window and unused admission status.'),
 ('7 Refunds and event cancellation recovery','Check attendee eligibility, record administrator review, process provider refunds and retain an audit trail.'),
 ('8 Notifications reports and administration','Deliver account and event updates; present sales and attendance reports; manage accounts and investigate recovery jobs.')],[1.8,5.0],9.5)
h('User roles and authorization',2)
table(['Role','Allowed actions','Scope restriction'],[
 ('Visitor','Browse published events and available tiers; register or sign in; use public discovery assistance.','No private orders tickets reports or organizer drafts.'),
 ('ATTENDEE','Manage own profile, save events, join alerts, reserve and pay, view private tickets and request refunds.','Every resource request is constrained to the authenticated attendee.'),
 ('ORGANIZER','Create and manage events and tiers, publish, cancel, inspect event orders and reports, and check tickets in.','Must own the event. Organizer role alone does not grant access to another organizer event.'),
 ('ADMIN','Manage user roles and status, inspect platform records, cancel events with an audit reason, review refunds and retry eligible recovery work.','Sensitive actions are audited; administrator access does not make an unverified payment successful.')],[1.0,3.5,2.3],9)
p('Visitor is an unauthenticated actor rather than a stored UserRole. Public self registration creates ATTENDEE accounts; callers cannot assign themselves privileged roles.')

doc.add_page_break();h('2 Basic backend architecture')
p('The application uses a modular monolith so related business rules can share explicit database transactions. Domain modules separate authentication, events, ticket tiers, orders, payments, tickets, refunds, discovery, notifications and reporting. The service layer uses Prisma directly; a separate repository layer is not required to explain the existing architecture.')
table(['Layer','Responsibility and boundaries'],[
 ('Client','Collect user intent and render confirmed state. Client side validation improves usability but does not authorize operations or establish stock, payment or admission outcomes.'),
 ('Browser BFF','Receive same origin browser requests; resolve an HttpOnly session cookie in Redis; attach encrypted server held tokens to API calls; enforce origin and CSRF checks on mutations. Direct API clients use Bearer access tokens.'),
 ('Express entry and middleware','Assign request IDs, parse bounded payloads, apply security headers, CORS and rate limits. Authenticate JWTs against current user status and token version; authorize allowed roles.'),
 ('Routes and schema validation','Match method and path; run Zod validation for path parameters, query strings and request bodies before invoking a controller.'),
 ('Controllers','Extract validated input and trusted identity, invoke services, and select the documented HTTP status and response envelope. Keep business decisions in services.'),
 ('Domain services','Enforce resource ownership, sales and refund windows, lifecycle rules, idempotency and financial invariants. Choose row locking and transaction boundaries.'),
 ('Prisma ORM and PostgreSQL','Execute queries and transactions; enforce PKs, FKs, unique keys and persisted state. PostgreSQL owns inventory, payments, passes and audit records.'),
 ('Jobs and integrations','Process leased durable jobs for expiry, provider verification, refunds and notifications. Provider network calls occur outside database transactions; validated results are applied in new transactions.'),
 ('Central error handler','Translate expected domain failures into safe JSON errors, log unexpected faults with a request ID, and prevent credentials or provider diagnostics from leaking.')],[1.45,5.35],9.4)
p('Redis is used for browser session infrastructure and optional public or AI caching. A cache is never used as the final authority for reservation inventory, money, refund eligibility or admission. If optional discovery caching fails, the API can read PostgreSQL; failure of required session storage prevents authentication from being assumed successful.')
section(True);diagram('eventgate-workflows-2.png','Figure 1 Backend layers and supporting services. A direct API consumer may call Express without the browser BFF.')
section(False)
h('3 SDA workflow')
h('Major feature Reserve tickets and confirm payment',2)
p('The workflow starts when an attendee reserves ticket inventory and ends when a verified payment results in issued passes. The reservation response is intentionally separate from the later payment outcome.')
numbered([
 'Client request. The attendee selects a tier and quantity. The browser BFF sends POST /api/v1/orders with a Bearer token and a stable Idempotency-Key. No client submitted price is trusted.',
 'Route and middleware. The API assigns a request ID, applies a rate limit, authenticates the active attendee, checks the required role, and validates the UUID and integer quantity from 1 to 10.',
 'Controller. The order controller reads the idempotency header and forwards the validated body and authenticated attendee ID to the order service.',
 'Business validation. The service checks the stored idempotency key and request hash, locks the related event and tier, and rechecks event publication, sales pause, sales window and available stock.',
 'Database transaction. The service reserves inventory, creates the order with price snapshots and an expiry 15 minutes later, and records idempotency and audit state. The transaction commits together or rolls back together. Serialization conflicts have bounded retries.',
 'API response. After a commit, the controller returns HTTP 201 and the confirmed order. Invalid input receives 400; unauthenticated access 401; disallowed roles 403; hidden or absent resources 404; unavailable stock or conflicting state 409.',
 'Checkout request. The attendee sends POST /payments/orders/:orderId/checkout. The service checks ownership, active reservation and checkout profile, persists a payment attempt, then calls SSLCommerz outside the transaction.',
 'Provider callback. A success, failure, cancellation or IPN payload is persisted with a deduplication key. Browser redirect parameters are not trusted as payment proof. A worker queries the provider validation API and compares transaction identity, amount and currency.',
 'Verified fulfillment. A new database transaction converts reserved inventory to sold inventory, marks the canonical purchase paid and issues one ticket for each purchased unit. Replayed callbacks do not issue additional passes.',
 'Readback and recovery. The attendee reads the authoritative order or attempt status and retrieves private QR passes when eligible. Expired reservations release stock. Late captures and ambiguous outcomes are reconciled, refunded or flagged for investigation rather than silently fulfilled.'
])
section(True);diagram('eventgate-workflows-1.png','Figure 2 Ticket reservation with atomic inventory updates and safe error responses. Payment verification follows as described in steps 7 to 10.')
section(False)
h('4 Database design')
p('PostgreSQL stores 23 entities. The ERD is divided into six named Draw.io views for readability. Tables repeated between views are references to the same physical table. The views show important fields; the dictionary and relationship ledger below record keys and cross view dependencies.')
link('ERD share link',links['erd'])
label('Notation','PK means primary key, FK means enforced foreign key, UQ means unique, and ? means nullable. Most identifiers are UUIDs. MaintenanceState.id is a text key. A parent with 1 to 0..N children is a one to many association. A unique child FK limits it to zero or one child per parent. Nullable FKs mean a child may have no parent for that association. Each connector is labelled with its relationship role and a : b, meaning a parents per child and b children per parent. For example 1 : 0..N is mandatory one parent per child with zero or many children per parent; 0..1 : 0..N permits a child without that parent. The canonical payment fulfillment link is optional on both sides, 0..1 : 0..1. The PaymentAttempt alias shows a self reference to the same physical table, not another entity.')
h('Relationships and integrity rules',2)
bullets([
 'One to many: an organizer User owns Events; an Event contains TicketTiers; an attendee User places Orders; an Order has PaymentAttempts and Tickets; a PaymentAttempt receives PaymentEvents.',
 'One to one where applicable: NotificationPreference.userId and EventPreference.userId are both PK and FK, so a User has zero or one row of each. Order.fulfilledPaymentAttemptId is nullable and unique, so a purchase selects at most one canonical attempt, and an attempt fulfills at most one order.',
 'Many to many: SavedEvent joins User and Event, unique on userId plus eventId. WaitlistEntry joins User and TicketTier, unique on userId plus ticketTierId. Orders also capture repeat purchase relationships rather than acting as a unique event membership.',
 'A ticket has a unique QR token and a unique orderId plus sequence. A payment attempt has a unique merchant transaction identifier and unique provider transaction fields when present. A callback event has a unique providerEventKey.',
 'Financial and ticket relationships generally use restrictive deletes to preserve history. PaymentEvent.paymentAttemptId uses SET NULL; EventPreference uses CASCADE on user deletion. Event and tier removal uses lifecycle rules and soft deletion where needed.',
 'Cross record rules such as a ticket tier belonging to the order event and the fulfillment attempt belonging to the order are enforced in services and transactions. A simple FK alone does not prove these business relationships.',
 'Scalar fields such as AuditLog.entityId, IdempotencyRecord.resourceId, PaymentEvent.capturedPaymentAttemptId and RefundActivity.actorId are not drawn as enforced FKs when the Prisma schema declares no relation.'
])
for i,caption in enumerate(['Commerce inventory and many to many memberships','Payments capture replay canonical fulfillment and refund settlement','Identity recovery and one to one preferences','Notifications jobs and audit operations','Optional delivery links for password recovery and email verification','Ticket ownership event inventory and optional admission staff'],1):
 section(True);diagram(f'eventgate-erd-{i}.png',f'Figure {i+2} {caption}. Refer to the foreign key ledger for associations that span views.')
section(False)
h('Entity dictionary',2)
descriptions={
 'User':'Identity contact details role active status and credential version.', 'Event':'Organizer owned event facts and publication lifecycle.', 'TicketTier':'Priced inventory and sale window for an event.', 'Order':'Attendee reservation with immutable event tier and price snapshots.', 'PaymentAttempt':'Provider checkout attempt with sequence and capture identity.', 'PaymentEvent':'Deduplicated provider callback receipt and processing state.', 'Ticket':'One admission pass per purchased unit with a unique opaque token.', 'Refund':'Provider linked refund request review and settlement lifecycle.', 'RefundActivity':'Append only history of refund actions.', 'AuditLog':'Actor and entity history for sensitive business operations.', 'IdempotencyRecord':'Request hash and remembered command outcome per scope and key.', 'Job':'Durable asynchronous work with run time lease and retry limits.', 'MaintenanceState':'Coordinator heartbeat and maintenance run summary.', 'RefreshSession':'Hashed refresh token session family and revocation state.', 'AuthIdentity':'External identity provider subject linked to a user.', 'PasswordReset':'Hashed single use password recovery token and expiry.', 'EmailVerification':'Hashed single use email verification token and expiry.', 'SavedEvent':'User event membership with reminder settings.', 'WaitlistEntry':'User tier stock alert membership and notification generation.', 'Notification':'In app message and source deduplication key.', 'NotificationPreference':'One user notification settings record.', 'NotificationDelivery':'Per channel message delivery retry and provider state.', 'EventPreference':'One user personalization settings record.'}
rows=[]
for name,data in models.items():
 pk=', '.join(f['name'] for f in data['fields'] if 'PK' in f['flags']);fks=', '.join(f['name'] for f in data['fields'] if 'FK' in f['flags']) or 'None'
 unique=['+'.join(f[1].split(', ')) for f in data['constraints'] if f[0]=='unique']+[f['name'] for f in data['fields'] if 'UQ' in f['flags']]
 rows.append((name,pk,fks,descriptions[name]+(' Unique: '+', '.join(unique)+'.' if unique else '')))
table(['Entity','PK','Enforced FK fields','Purpose and important unique constraints'],rows,[1.3,.55,1.65,3.3],8.3)
h('Complete foreign key ledger',2)
relrows=[]
for name,data in models.items():
 for f in data['fields']:
  if f['target']:
   card=('0..1' if '?' in f['type'] else '1')+' parent per child; '+('0..1' if ('UQ' in f['flags'] or 'PK' in f['flags']) else '0..N')+' children per parent'
   relrows.append((name+'.'+f['name'],f['target']+'.id',card))
table(['Child FK','Referenced PK','Cardinality'],relrows,[2.85,1.55,2.4],8.5)
h('Storage and indexing decisions',2)
p('All amounts are integer paisa with BDT currency. Store timestamps in UTC using timezone aware columns and present event times in Asia/Dhaka. Preserve order snapshots so later event or tier edits do not rewrite purchase history. Audit metadata and provider payloads must be redacted before leaving authorized operational views.')
p('Indexes follow access paths: Event organizer plus lifecycle and start time; Order attendee plus created time, event plus status, and status plus reservation expiry; Ticket event plus status and check in time; User role plus status; Job status plus runAt and lease expiry; and notification user plus readAt or createdAt. Unique keys provide the final duplicate barrier in addition to application checks.')

h('5 API design conventions')
label('Base path','/api/v1. The catalogue below lists paths relative to that base.')
label('Authentication','Protected API requests use Authorization: Bearer <accessToken>. Authentication validates the current account and token version. Role checks are followed by ownership checks. Browser clients use the BFF session cookie and a CSRF header on mutations; JWTs are not kept in browser local storage.')
label('Payloads','Use application/json except provider callbacks, which accept bounded form encoded bodies. UUID path parameters are required where indicated by names such as eventId or orderId. Reject unknown fields in strict command schemas.')
label('Idempotency','POST /orders and POST /payments/orders/:orderId/checkout require Idempotency-Key with 1 to 255 characters. Retry the same key and same command after a network failure. Reusing the key for different input returns a conflict. Financial recovery endpoints do not manufacture successful provider outcomes.')
label('Pagination','List operations return data plus meta containing page, limit, total and totalPages. Common defaults are page 1 and limit 20 with maximum 100; endpoints may choose their own defaults. Public discovery supports search, city, category, date, sort and budget filters. Dates use ISO 8601 values.')
label('Response rules','Success uses success, message and data. Errors use success false, message and an errors array with code, message and optional field. A request ID header supports diagnosis. Examples below contain illustrative UUIDs and values and show essential fields rather than every optional server projection.')
code(ok('Example response',{'id':OID,'status':'PENDING_PAYMENT'}))
code(err('Validation failed','INVALID_VALUE','Quantity must be an integer between 1 and 10.','body.quantity'))
table(['Status','Meaning'],[('200 or 201','Confirmed read or synchronous command; 201 creates a resource.'),('202','Accepted or pending verification; it does not confirm payment or refund settlement.'),('400','Invalid JSON schema field parameters or required command headers.'),('401 or 403','Missing invalid or revoked authentication; insufficient role or inactive account.'),('404','Missing resource or a resource concealed by ownership rules.'),('409','Conflicting stock lifecycle idempotency or refund eligibility.'),('413 or 429','Payload limit or rate limit exceeded.'),('500 or 503','Unexpected server failure or an unavailable dependency.')],[1.0,5.8],9.3)
p('Provider IPN acknowledgement is plain text OK. Browser payment callbacks return HTTP 303 to the application; neither response is proof of settlement. Check in business outcomes can use HTTP 200 with a resultCode such as ACCEPTED or ALREADY_CHECKED_IN, so the client must inspect the resultCode.')

section(True)
h('Complete API operation catalogue')
p('All paths below use /api/v1. Bearer requires a current active user. Ownership restrictions are enforced in services as well as role middleware. Public AI and demo endpoints are configuration dependent. The catalogue covers all 109 API route operations in the design basis.')
groups=defaultdict(list)
for a in api:
 ep=a['endpoint']
 key='Administration and operations' if ep.startswith('/admin') else 'Organizer events and reports' if ep.startswith('/organizer') else 'Authentication and account' if ep.startswith(('/auth','/users')) else 'Public discovery and assistance' if ep.startswith(('/events','/ai','/health')) and a['auth']=='No' else 'Attendee purchases tickets and engagement'
 groups[key].append(a)
for group,items in groups.items():
 h(group,2)
 table(['Method','Endpoint','Purpose','Auth','Role or scope'],[(a['method'],a['endpoint'],a['purpose'],a['auth'],a['role'].replace('ORGANIZER owns event; ADMIN may access','ORGANIZER owner or ADMIN').replace('ATTENDEE; own resources only','ATTENDEE own resources').replace('Any active user; own account only','Active user own account').replace('Public; rate limited where applicable','Public')) for a in items],[.6,3.7,2.5,1.0,2.6],8.8)
section(False)
h('6 Detailed design of ten important APIs')
p('Each contract below includes input, authentication, successful output, failure behavior and business rules. Access and refresh tokens in examples are placeholders, never working credentials. Future event dates keep the examples valid after the design date. Response examples are compact projections of the safe resource fields.')
publicuser={'id':UID,'email':'attendee@example.com','displayName':'Ayesha Rahman','role':'ATTENDEE','isQuickAccess':False}
authdata={'user':publicuser,'accessToken':'<access-token>','refreshToken':'<opaque-refresh-token>','accessTokenExpiresIn':900}
event={'id':EID,'slug':'dhaka-tech-forum-2027','title':'Dhaka Tech Forum 2027','description':'A technology event with talks and workshops for the local developer community.','category':'Technology','venue':'Convention Hall','city':'Dhaka','address':'Agargaon Dhaka Bangladesh','startAt':'2027-01-20T09:00:00.000Z','endAt':'2027-01-20T15:00:00.000Z','imageUrl':None,'status':'DRAFT'}
tier={'id':TID,'eventId':EID,'name':'General Admission','pricePaisa':50000,'capacity':100,'reservedQuantity':0,'soldQuantity':0,'salesStartAt':'2027-01-01T00:00:00.000Z','salesEndAt':'2027-01-20T08:00:00.000Z','salesPaused':False}
order={'id':OID,'eventId':EID,'ticketTierId':TID,'eventNameSnapshot':event['title'],'ticketTierNameSnapshot':'General Admission','quantity':2,'unitPricePaisaSnapshot':50000,'totalAmountPaisaSnapshot':100000,'currency':'BDT','status':'PENDING_PAYMENT','reservationExpiresAt':'2027-01-10T10:15:00.000Z'}
contracts=[
 ('01 Register attendee','POST','/auth/register','Public; no access token. Rate limited.','No path or query parameters.',{'displayName':'Ayesha Rahman','email':'attendee@example.com','password':'ExamplePassword2027!'},201,ok('Account registered',authdata),409,err('Email is already registered','EMAIL_ALREADY_REGISTERED','Use another email.','body.email'),'Normalize email and validate displayName from 2 to 120 characters and password from 12 to 128 characters. Hash passwords with Argon2id. Force ATTENDEE role and create the refresh session atomically. Return 400 for invalid fields and 429 for excessive requests.'),
 ('02 Sign in','POST','/auth/login','Public; no access token. Rate limited.','No path or query parameters.',{'email':'attendee@example.com','password':'ExamplePassword2027!'},200,ok('Signed in',authdata),401,err('Invalid email or password','INVALID_CREDENTIALS','Email or password is incorrect.'),'Verify the password hash, active account state and credential version before issuing tokens. Use the same generic credential error for unknown email and incorrect password. Inactive accounts cannot sign in. The BFF stores tokens encrypted in its server side session.'),
 ('03 Browse published events','GET','/events','Public. No personal session is required.','Query: page=1, limit=12, city=Dhaka, category=Technology, sort=soonest. Optional search, startAfter, startBefore and maxPricePaisa.',None,200,{'success':True,'message':'Events retrieved','data':[{'id':EID,'slug':event['slug'],'title':event['title'],'status':'PUBLISHED','startAt':event['startAt'],'endAt':event['endAt'],'venue':event['venue'],'city':'Dhaka','category':'Technology','imageUrl':None,'organizer':{'displayName':'Dhaka Events'},'ticketTiers':[{'id':TID,'name':'General Admission','pricePaisa':50000,'availableQuantity':100,'salesPaused':False,'salesStartAt':tier['salesStartAt'],'salesEndAt':tier['salesEndAt'],'capacity':100}]}],'meta':{'page':1,'limit':12,'total':1,'totalPages':1}},400,err('Validation failed','INVALID_VALUE','Page must be a positive integer.','query.page'),'Expose only published non deleted listings and return stable pagination. Apply budget matching to currently bookable tier inventory. Discovery facts may be cached; GET /events/:slug/availability and order creation use fresh eligibility and stock checks. Private drafts are excluded.'),
 ('04 Create an event draft','POST','/organizer/events','Bearer; ORGANIZER or ADMIN. Owner is derived from the authenticated identity.','No path parameters.',{k:v for k,v in event.items() if k not in ['id','status']},201,ok('Event created',event),400,err('Validation failed','INVALID_VALUE','End time must be after start time.','body.endAt'),'Require a future start and an end after start. Validate title, category, location and optional HTTPS cover URL. Enforce a unique normalized slug. The client cannot assign organizerId or publish status. A created draft must go through publication readiness checks.'),
 ('05 Add ticket inventory','POST','/organizer/events/:eventId/ticket-tiers','Bearer; event owning ORGANIZER or ADMIN.','Path eventId: '+EID,{k:v for k,v in tier.items() if k not in ['id','eventId','reservedQuantity','soldQuantity']},201,ok('Ticket tier created',tier),400,err('Validation failed','INVALID_VALUE','Price must be a positive integer in paisa.','body.pricePaisa'),'Validate positive integer pricePaisa and capacity and a valid sales window within event rules. The service sets reserved and sold counts. Enforce unique eventId plus name. Editing capacity must not reduce it below sold plus reserved inventory. Missing or foreign events return 404.'),
 ('06 Publish an event','POST','/organizer/events/:eventId/publish','Bearer; event owning ORGANIZER or ADMIN.','Path eventId: '+EID,{},200,ok('Event published',{**event,'status':'PUBLISHED','publishedAt':'2027-01-01T08:00:00.000Z'}),409,err('Event is not ready to publish','NO_ELIGIBLE_TIER','Add an eligible unpaused ticket tier before publication.'),'Revalidate the saved draft, future schedule and eligible saved tier in the publication transaction. PUBLICATION_NOT_READY is the service failure category and blocker details identify the cause. Do not publish from client supplied preview state. Invalidate public discovery facts after a successful change.'),
 ('07 Reserve tickets','POST','/orders','Bearer; ATTENDEE. Header Idempotency-Key: reservation-example-001.','No path parameters. Example request time is 2027-01-10T10:00:00Z.',{'ticketTierId':TID,'quantity':2},201,ok('Order created and inventory reserved',order),409,err('Not enough tickets are available','INSUFFICIENT_INVENTORY','Choose a smaller quantity or another tier.'),'Require quantity from 1 to 10. Lock and recheck current sale state and stock; use a serializable transaction to increase reservedQuantity and create the order, snapshots, audit record and idempotency result together. Same key plus same payload recovers the same order; different payload returns IDEMPOTENCY_KEY_REUSED. A 15 minute expiry is evaluated by server time.'),
 ('08 Start or recover checkout','POST','/payments/orders/:orderId/checkout','Bearer; order owning ATTENDEE. Header Idempotency-Key: checkout-example-001.','Path orderId: '+OID,{},200,ok('SSLCommerz checkout available',{'status':'AVAILABLE','paymentAttemptId':PID,'merchantTransactionId':'EG-EXAMPLE-001','checkoutUrl':'https://sandbox.sslcommerz.com/EasyCheckOut/example-session','checkoutExpiresAt':'2027-01-10T10:15:00.000Z'}),409,err('Order is not payable','ORDER_NOT_PAYABLE','Reservation is expired cancelled or already paid.'),'Validate order ownership, an unexpired active hold, current sale eligibility and required contact data. Persist the attempt before the provider request. A 202 PROCESSING or UNKNOWN response includes paymentAttemptId, paymentStatus and statusPath for readback; it contains no usable checkout URL. Retrying the same key must not start a duplicate financial attempt.'),
 ('09 Evaluate a QR admission','POST','/organizer/events/:eventId/check-ins','Bearer; event owning ORGANIZER or ADMIN.','Path eventId: '+EID,{'qrToken':'A'*43},200,ok('Check-in evaluated',{'resultCode':'ACCEPTED','id':TKID,'status':'CHECKED_IN','ticket':{'id':TKID,'sequence':1,'checkedInAt':'2027-01-20T09:10:00.000Z','ticketTier':{'name':'General Admission'},'attendee':{'displayName':'Ayesha Rahman'}}}),400,err('Validation failed','INVALID_VALUE','Provide the opaque EventGate QR token not a URL.','body.qrToken'),'Validate the 43 character opaque token. Check event ownership, the admission window from two hours before start through end, verified purchase and unused ticket state. Lock the order graph, atomically update the pass and write an audit record. Business rejections also use 200 with resultCode INVALID_TICKET, WRONG_EVENT, ALREADY_CHECKED_IN, WINDOW_CLOSED or ADMISSION_BLOCKED. A 200 envelope alone never admits a person.'),
 ('10 Request an attendee refund','POST','/orders/:orderId/refund-requests','Bearer; order owning ATTENDEE.','Path orderId: '+OID,{'userMessage':'I can no longer attend this event.'},201,ok('Refund requested',{'id':RID,'orderId':OID,'paymentAttemptId':PID,'reason':'ATTENDEE_REQUEST','amountPaisa':100000,'currency':'BDT','status':'REQUESTED','requestedAt':'2027-01-12T10:00:00.000Z','userMessage':'I can no longer attend this event.'}),409,err('This booking is not eligible for an attendee refund.','REFUND_WINDOW_CLOSED','Refund requests require at least 24 hours before the event starts.'),'Require a paid canonical verified purchase, a published event, at least 24 hours before start and all tickets unused and ACTIVE. Reject conflicting live refunds or investigation states. Lock related financial records and create the request with audit activity. Amount and currency come from the stored purchase; the client cannot choose either. REQUESTED does not mean money has been returned.')
]
assert len(contracts)==10
for i,c in enumerate(contracts):

 if i:doc.add_page_break()
 title,method,path,auth,params,body,status,response,error_status,error_response,rules=c
 h(title,2);p(method+' /api/v1'+path,'Heading 3')
 label('Authentication',auth);label('Parameters',params)
 if body is not None:
  h('Request body',3);code(body)
 else:label('Request body','None. Filters are query parameters.')
 h(f'Successful response HTTP {status}',3);code(response)
 h(f'Possible error response HTTP {error_status}',3);code(error_response)
 h('Validation and business behavior',3);p(rules)
(ROOT/'design/detailed-api-contracts.json').write_text(json.dumps([{'title':c[0],'method':c[1],'endpoint':c[2],'authentication':c[3],'parameters':c[4],'request':c[5],'successStatus':c[6],'success':c[7],'errorStatus':c[8],'error':c[9],'rules':c[10]} for c in contracts],indent=2))

doc.add_page_break();h('7 State transitions and reliability')
table(['Resource','Lifecycle and controlling rule'],[
 ('Event','DRAFT to PUBLISHED after readiness; published events can become CANCELLED or COMPLETED. Dates and location are locked after publication where the domain requires it.'),
 ('Order','PENDING_PAYMENT to PAID only after a verified canonical capture. Unpaid reservations can become EXPIRED or CANCELLED and release stock once. A settled refund can move the purchase to REFUNDED.'),
 ('Payment attempt','INITIATING and PENDING progress to a validated terminal outcome. UNKNOWN requires reconciliation. Preserve every attempt; identify one canonical purchase and separately handle excess captures or replay attempts.'),
 ('Ticket','ACTIVE to CHECKED_IN once. Refund or cancellation can block or void admission. QR retrieval checks effective admission state; a cached image never overrides server check in state.'),
 ('Refund','REQUESTED is reviewed as APPROVED or REJECTED. Processing becomes SUCCEEDED only after provider confirmation; UNKNOWN and investigated failures require reconciliation before retry.'),
 ('Job','PENDING work is leased as RUNNING. Completed work is SUCCEEDED; recoverable failures are bounded and retried; expired leases permit safe recovery. Deduplication prevents duplicate scheduled business work.')],[1.25,5.55],9.2)
h('Concurrency failure and security controls',2)
bullets([
 'Inventory invariant: available = capacity minus reservedQuantity minus soldQuantity, with available never negative. Business writes lock the relevant tier and recheck the current event state in the transaction.',
 'Canonical fulfillment: a verified capture may fulfill its order only once. Unique provider identifiers, unique ticket order sequences and idempotency records make replay safe.',
 'Provider uncertainty: a timeout is not a failed charge and is not a successful charge. Persist the attempt, return a pending outcome and schedule readback. Late money is handled explicitly.',
 'Reservation expiry: release reserved stock once, preserve the order history, and trigger availability notifications without treating the notification as a ticket hold.',
 'Data privacy: owner constrained reads return only safe projections. Raw provider data, QR tokens and account recovery secrets are not included in public event responses or logs.',
 'Operational recovery: retain audit events, structured request IDs and bounded retry diagnostics. Use health probes, graceful shutdown and leased work so infrastructure restarts do not silently drop financial work.',
 'Performance: cache discovery facts and coalesce equivalent reads, paginate lists, avoid broad blocking refetches, index common access paths and keep provider calls outside locked transactions. Measure latency on the deployed environment; deployment placement is an operational decision rather than an assumed speed guarantee.'
])
h('Design verification checklist',2)
table(['Scenario','Expected invariant'],[
 ('Two attendees request the last remaining ticket','At most one reservation succeeds; no negative inventory.'),
 ('A client retries an order after a network timeout','The same key and payload recover one order without another stock hold.'),
 ('A payment callback is replayed','One canonical fulfillment and the original number of issued passes.'),
 ('The provider times out after taking payment','The attempt remains uncertain until verified; no blind duplicate checkout.'),
 ('Two staff scan the same QR simultaneously','Only one admission is ACCEPTED; later evaluation reports already used.'),
 ('An organizer requests another organizer event','The service denies access without disclosing private details.'),
 ('A refund races with check in','Financial locks and eligibility rules prevent inconsistent admission and settlement.'),
 ('A job worker restarts during processing','Lease recovery and deduplication allow bounded safe reprocessing.')],[2.8,4.0],9.2)
h('Design basis and references',2)
p('The database names, constraints, route inventory and business rules were reviewed against the EventGate backend source at revision f4f9396e83c317bc8b71e5a09762a1cd148ada9e. The examples specify expected API behavior with illustrative data; this document is a design deliverable and does not assert that these example requests were executed against production.')
link('Source Prisma database schema','https://github.com/SayemHasan74/eventgate-api/blob/f4f9396e83c317bc8b71e5a09762a1cd148ada9e/prisma/schema.prisma')
link('Source API route mounts','https://github.com/SayemHasan74/eventgate-api/blob/f4f9396e83c317bc8b71e5a09762a1cd148ada9e/src/app.ts')
link('Editable ERD and workflow sources',links['repository'])
file=OUT/'EventGate Backend System Design.docx';doc.save(file)
print(file)
print('API operations:',len(api),'entities:',len(models),'FK associations:',len(relrows),'detailed contracts:',len(contracts))
